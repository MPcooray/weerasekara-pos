import io
import os
from concurrent.futures import ThreadPoolExecutor
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routes.product_imports import csv_safe
from app.core.security import hash_password
from app.db.session import SessionLocal
from app.models.inventory import InventoryBalance, StockMovement
from app.models.product import Product
from app.models.supplier import Supplier
from app.models.tenant import Tenant
from app.models.user import User, UserRole


def workbook(client, rows, version="1.0", formula=False):
    response = client.get("/api/v1/product-imports/template")
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    book["Instructions"]["B2"] = version
    sheet = book["Products"]
    sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append(row)
    if formula:
        sheet["J2"] = "=1+1"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def row(
    sku="000123",
    barcode="04791234567890",
    name="Imported Frame",
    category="spectacle_frame",
    supplier="",
    cost=4500,
    price=7900,
    quantity=5,
    reorder=2,
    track="Yes",
):
    return [
        sku,
        barcode,
        name,
        category,
        "Tetrova",
        "M1",
        "Black",
        "52-18",
        supplier,
        cost,
        price,
        quantity,
        reorder,
        "Imported fixture",
        track,
        "Acetate",
        "Rectangle",
        "",
        "",
        "No",
    ]


def preview(client, content, mode="create_only", filename="products.xlsx"):
    return client.post(
        "/api/v1/product-imports/preview",
        data={"mode": mode},
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def login_as(client, tenant_slug, email, role):
    with SessionLocal() as db:
        tenant = db.query(Tenant).filter_by(slug=tenant_slug).one()
        db.add(
            User(
                tenant_id=tenant.id,
                email=email,
                full_name=role.value.title(),
                password_hash=hash_password("SafePassword123!"),
                role=role,
            )
        )
        db.commit()
    result = client.post(
        "/api/v1/auth/login",
        json={"tenant_slug": tenant_slug, "email": email, "password": "SafePassword123!"},
    )
    return {"Authorization": f"Bearer {result.json()['access_token']}"}


def test_template_preview_confirm_idempotency_and_stock_ledger(authenticated_client):
    assert csv_safe("=malicious") == "'=malicious"
    template = authenticated_client.get("/api/v1/product-imports/template")
    assert template.status_code == 200
    book = load_workbook(io.BytesIO(template.content))
    assert book.sheetnames == ["Instructions", "Products"]
    assert book["Instructions"]["B2"].value == "1.0"
    assert book["Products"]["A2"].number_format == "@"
    assert book["Products"]["B2"].value == "04791234567890"

    result = preview(authenticated_client, workbook(authenticated_client, [row()]))
    assert result.status_code == 200, result.text
    data = result.json()
    assert (data["total_rows"], data["valid_rows"], data["new_products"]) == (1, 1, 1)
    with SessionLocal() as db:
        assert db.query(Product).count() == 0
        assert db.query(StockMovement).count() == 0

    confirmed = authenticated_client.post(
        "/api/v1/product-imports/confirm",
        json={"preview_token": data["preview_token"]},
        headers={"Idempotency-Key": "product-import-once"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["created_count"] == 1
    assert confirmed.json()["stock_movements_count"] == 1
    repeated = authenticated_client.post(
        "/api/v1/product-imports/confirm",
        json={"preview_token": data["preview_token"]},
        headers={"Idempotency-Key": "product-import-once"},
    )
    assert repeated.json()["created_count"] == 1
    with SessionLocal() as db:
        product = db.query(Product).one()
        assert product.sku == "000123"
        assert product.barcode == "04791234567890"
        assert db.query(StockMovement).one().balance_after == 5
        assert db.query(InventoryBalance).one().quantity == 5


def test_row_validation_duplicates_supplier_prices_and_quantities(authenticated_client):
    rows = [
        row(sku="DUP-1", barcode="00001111"),
        row(sku="DUP-1", barcode="00002222"),
        row(sku="BAD-CAT", barcode="00003333", category="unknown"),
        row(sku="BAD-SUP", barcode="00004444", supplier="Missing Supplier"),
        row(sku="BAD-PRICE", barcode="00005555", cost="12.345"),
        row(sku="BAD-QTY", barcode="00006666", quantity="1.5"),
        row(sku="BAD-STOCK", barcode="00007777", quantity=2, track="No"),
    ]
    result = preview(authenticated_client, workbook(authenticated_client, rows))
    assert result.status_code == 200
    data = result.json()
    assert data["invalid_rows"] == 6
    messages = " ".join(error for item in data["rows"] for error in item["errors"])
    assert "duplicate inside workbook" in messages
    assert "Category" in messages
    assert "no active supplier" in messages
    assert "decimal places" in messages
    assert "nonnegative integer" in messages
    assert "requires Track Stock" in messages
    blocked = authenticated_client.post(
        "/api/v1/product-imports/confirm",
        json={"preview_token": data["preview_token"]},
        headers={"Idempotency-Key": "invalid-import"},
    )
    assert blocked.status_code == 422


def test_existing_modes_update_fields_but_never_stock(authenticated_client):
    created = authenticated_client.post(
        "/api/v1/products",
        json={
            "name": "Old",
            "sku": "EXIST-1",
            "barcode": "00009999",
            "category": "spectacle_frame",
            "cost_price_minor": 100,
            "selling_price_minor": 200,
            "track_stock": True,
            "reorder_threshold": 1,
        },
    )
    assert created.status_code == 201
    content = workbook(
        authenticated_client, [row(sku="EXIST-1", barcode="00009999", name="Updated", quantity=99)]
    )
    create_only = preview(authenticated_client, content)
    assert create_only.json()["invalid_rows"] == 1
    update = preview(authenticated_client, content, "update_existing")
    assert update.json()["valid_rows"] == 1
    assert "ignored" in update.json()["rows"][0]["warnings"][0]
    confirmed = authenticated_client.post(
        "/api/v1/product-imports/confirm",
        json={"preview_token": update.json()["preview_token"]},
        headers={"Idempotency-Key": "update-existing-product"},
    )
    assert confirmed.status_code == 200
    with SessionLocal() as db:
        assert db.query(Product).one().name == "Updated"
        assert db.query(StockMovement).count() == 0


def test_malformed_formula_version_extension_and_size_rejected(authenticated_client):
    assert preview(authenticated_client, b"not-a-workbook").status_code == 422
    assert preview(authenticated_client, b"x", filename="products.xls").status_code == 415
    assert preview(authenticated_client, b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    wrong = preview(authenticated_client, workbook(authenticated_client, [row()], version="99"))
    assert wrong.status_code == 422
    formulas = preview(authenticated_client, workbook(authenticated_client, [row()], formula=True))
    assert formulas.status_code == 422


def test_permissions_tenant_binding_and_atomic_supplier_change(authenticated_client, client):
    supplier = authenticated_client.post(
        "/api/v1/suppliers", json={"business_name": "Import Supplier"}
    ).json()
    data = preview(
        authenticated_client,
        workbook(
            authenticated_client,
            [
                row(sku="ATOMIC-1", barcode="00008881"),
                row(sku="ATOMIC-2", barcode="00008882", supplier="Import Supplier"),
            ],
        ),
    ).json()
    cashier = login_as(client, "weerasekara-optics", "import-cashier@example.com", UserRole.CASHIER)
    assert client.get("/api/v1/product-imports/template", headers=cashier).status_code == 403

    with SessionLocal() as db:
        other = Tenant(name="Other Import Tenant", slug="other-import-tenant")
        db.add(other)
        db.flush()
        manager = User(
            tenant_id=other.id,
            email="manager@other-import.example",
            full_name="Other Manager",
            password_hash=hash_password("SafePassword123!"),
            role=UserRole.MANAGER,
        )
        db.add(manager)
        supplier_row = db.get(Supplier, UUID(supplier["id"]))
        supplier_row.is_archived = True
        db.commit()
    login = client.post(
        "/api/v1/auth/login",
        json={
            "tenant_slug": "other-import-tenant",
            "email": "manager@other-import.example",
            "password": "SafePassword123!",
        },
    )
    other_headers = {
        "Authorization": f"Bearer {login.json()['access_token']}",
        "Idempotency-Key": "cross-tenant-import",
    }
    assert (
        client.post(
            "/api/v1/product-imports/confirm",
            json={"preview_token": data["preview_token"]},
            headers=other_headers,
        ).status_code
        == 404
    )

    owner = client.post(
        "/api/v1/auth/login",
        json={
            "tenant_slug": "weerasekara-optics",
            "email": "staff@example.com",
            "password": "SafePassword123!",
        },
    )
    owner_headers = {
        "Authorization": f"Bearer {owner.json()['access_token']}",
        "Idempotency-Key": "atomic-supplier-change",
    }
    failed = client.post(
        "/api/v1/product-imports/confirm",
        json={"preview_token": data["preview_token"]},
        headers=owner_headers,
    )
    assert failed.status_code == 409
    with SessionLocal() as db:
        assert db.query(Product).filter(Product.sku.in_(["ATOMIC-1", "ATOMIC-2"])).count() == 0


@pytest.mark.skipif(
    not os.environ.get("DATABASE_URL", "sqlite").startswith("postgresql"),
    reason="Concurrent uniqueness protection is verified on PostgreSQL",
)
def test_concurrent_imports_create_one_product(authenticated_client):
    first = preview(
        authenticated_client,
        workbook(
            authenticated_client,
            [row(sku="RACE-1", barcode="00007891", quantity=3)],
        ),
    ).json()
    second = preview(
        authenticated_client,
        workbook(
            authenticated_client,
            [row(sku="RACE-1", barcode="00007891", quantity=3)],
        ),
    ).json()
    authorization = authenticated_client.headers["Authorization"]

    def submit(token, key):
        with TestClient(authenticated_client.app) as contender:
            return contender.post(
                "/api/v1/product-imports/confirm",
                json={"preview_token": token},
                headers={"Authorization": authorization, "Idempotency-Key": key},
            ).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = sorted(
            pool.map(
                lambda args: submit(*args),
                [
                    (first["preview_token"], "concurrent-import-one"),
                    (second["preview_token"], "concurrent-import-two"),
                ],
            )
        )
    assert statuses == [200, 409]
    with SessionLocal() as db:
        product = db.query(Product).filter_by(sku="RACE-1").one()
        assert db.query(StockMovement).filter_by(product_id=product.id).count() == 1
