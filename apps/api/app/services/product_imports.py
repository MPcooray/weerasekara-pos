import hashlib
import io
import secrets
import zipfile
from datetime import UTC, datetime, timedelta
from decimal import Decimal, InvalidOperation
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.inventory import MovementType
from app.models.product import Product, ProductCategory
from app.models.product_import import ProductImport, ProductImportMode, ProductImportStatus
from app.models.supplier import Supplier
from app.models.user import User
from app.schemas.inventory import ProductCreate, StockMutation
from app.services.audit import record_audit
from app.services.inventory import mutate_stock

TEMPLATE_VERSION = "1.0"
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 30
HEADERS = [
    "SKU*",
    "Barcode",
    "Product Name*",
    "Category*",
    "Brand",
    "Model",
    "Color",
    "Size",
    "Supplier",
    "Cost Price (LKR)*",
    "Selling Price (LKR)*",
    "Initial Quantity*",
    "Reorder Level*",
    "Description",
    "Track Stock*",
    "Frame Material",
    "Frame Shape",
    "Lens Type",
    "Lens Material",
    "Made To Order",
]


def template_bytes() -> bytes:
    book = Workbook()
    info = book.active
    info.title = "Instructions"
    info.append(["Tetrova Optical POS Product Import", ""])
    info.append(["Template Version", TEMPLATE_VERSION])
    info.append(
        [
            "Instructions",
            "Do not rename sheets or columns. Required columns end in *. "
            "Use existing supplier names exactly. Prices are LKR. Initial quantity "
            "applies only to newly created tracked products.",
        ]
    )
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    sheet = book.create_sheet("Products")
    sheet.append(HEADERS)
    sheet.append(
        [
            "FRM-0001",
            "04791234567890",
            "Classic Acetate Frame",
            "spectacle_frame",
            "Tetrova",
            "A100",
            "Black",
            "52-18",
            "",
            4500,
            7900,
            10,
            2,
            "Sample row—replace or remove",
            "Yes",
            "Acetate",
            "Rectangle",
            "",
            "",
            "No",
        ]
    )
    sheet.append(
        [
            "SRV-EYE",
            "",
            "Comprehensive Eye Examination",
            "optical_service",
            "",
            "",
            "",
            "",
            "",
            0,
            1500,
            0,
            0,
            "Sample service",
            "No",
            "",
            "",
            "",
            "",
            "No",
        ]
    )
    fill = PatternFill("solid", fgColor="183154")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:T{MAX_ROWS + 1}"
    for column in ("A", "B"):
        for row in range(2, MAX_ROWS + 2):
            sheet[f"{column}{row}"].number_format = "@"
    for column in ("J", "K"):
        for row in range(2, MAX_ROWS + 2):
            sheet[f"{column}{row}"].number_format = "0.00"
    widths = [18, 20, 30, 22, 18, 18, 14, 14, 26, 18, 20, 17, 16, 36, 15, 18, 18, 18, 18, 16]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    categories = '"' + ",".join(item.value for item in ProductCategory) + '"'
    sheet.add_data_validation(DataValidation(type="list", formula1=categories, allow_blank=False))
    sheet.data_validations.dataValidation[-1].add(f"D2:D{MAX_ROWS + 1}")
    for column in ("O", "T"):
        validation = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{MAX_ROWS + 1}")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def _safe_workbook(content: bytes):
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(413, "Workbook exceeds the 5 MB limit")
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise HTTPException(422, "The uploaded file is not a valid .xlsx workbook")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > 200:
            raise HTTPException(422, "Workbook contains too many internal files")
        if sum(item.file_size for item in infos) > 25 * 1024 * 1024:
            raise HTTPException(422, "Workbook expanded size is unsafe")
        if any(item.file_size > 10 * 1024 * 1024 for item in infos):
            raise HTTPException(422, "Workbook contains an oversized internal file")
        if any("externalLink" in item.filename for item in infos):
            raise HTTPException(422, "External workbook links are not supported")
    try:
        book = load_workbook(
            io.BytesIO(content), read_only=False, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise HTTPException(422, "The workbook could not be read safely") from exc
    if getattr(book, "_external_links", []):
        raise HTTPException(422, "External workbook links are not supported")
    return book


def _text(value, limit: int, required: bool = False) -> str | None:
    result = "" if value is None else str(value).strip()
    if required and not result:
        raise ValueError("is required")
    if len(result) > limit:
        raise ValueError(f"must be at most {limit} characters")
    return result or None


def _money(value) -> int:
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError("must be a valid LKR amount") from None
    if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
        raise ValueError("must be nonnegative with at most two decimal places")
    minor = int(amount * 100)
    if minor > 2_000_000_000:
        raise ValueError("is too large")
    return minor


def _integer(value, field: str) -> int:
    try:
        decimal = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field} must be an integer") from None
    if decimal != decimal.to_integral_value() or decimal < 0 or decimal > 1_000_000:
        raise ValueError(f"{field} must be a nonnegative integer")
    return int(decimal)


def _boolean(value, default: bool) -> bool:
    if value is None or str(value).strip() == "":
        return default
    normalized = str(value).strip().lower()
    if normalized in {"yes", "true", "1"}:
        return True
    if normalized in {"no", "false", "0"}:
        return False
    raise ValueError("must be Yes or No")


def parse_preview(db: Session, user: User, content: bytes, mode: ProductImportMode) -> dict:
    book = _safe_workbook(content)
    if "Instructions" not in book.sheetnames or "Products" not in book.sheetnames:
        raise HTTPException(422, "Official Instructions and Products sheets are required")
    if str(book["Instructions"]["B2"].value).strip() != TEMPLATE_VERSION:
        raise HTTPException(422, f"Unsupported template version; expected {TEMPLATE_VERSION}")
    sheet = book["Products"]
    if sheet.max_column > MAX_COLUMNS or sheet.max_row > MAX_ROWS + 1:
        raise HTTPException(422, f"Workbook limit is {MAX_ROWS} rows and {MAX_COLUMNS} columns")
    actual = [cell.value for cell in sheet[1]][: len(HEADERS)]
    if actual != HEADERS:
        raise HTTPException(422, "Products sheet columns do not match the official template")
    for row in sheet.iter_rows():
        if any(cell.data_type == "f" for cell in row):
            raise HTTPException(422, "Excel formulas are not accepted")
    suppliers = {
        item.business_name.casefold(): item
        for item in db.scalars(
            select(Supplier).where(
                Supplier.tenant_id == user.tenant_id, Supplier.is_archived.is_(False)
            )
        ).all()
    }
    existing = {
        item.sku: item
        for item in db.scalars(select(Product).where(Product.tenant_id == user.tenant_id)).all()
    }
    existing_barcodes = {item.barcode: item.sku for item in existing.values() if item.barcode}
    rows, seen_skus, seen_barcodes = [], set(), set()
    for number, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True), 2
    ):
        if not any(value not in (None, "") for value in values):
            continue
        errors, warnings = [], []

        def capture(label, function, fallback=None, errors=errors):
            try:
                return function()
            except (ValueError, ValidationError) as exc:
                errors.append(f"{label}: {str(exc).splitlines()[0]}")
                return fallback

        sku = capture("SKU", lambda values=values: _text(values[0], 64, True).upper(), "")
        barcode = capture("Barcode", lambda values=values: _text(values[1], 64), None)
        name = capture("Product Name", lambda values=values: _text(values[2], 180, True), "")
        category = capture(
            "Category",
            lambda values=values: ProductCategory(str(values[3]).strip()).value,
            "",
        )
        supplier_name = capture("Supplier", lambda values=values: _text(values[8], 160), None)
        supplier = suppliers.get(supplier_name.casefold()) if supplier_name else None
        if supplier_name and not supplier:
            errors.append("Supplier: no active supplier with this exact name")
        if sku and (sku in seen_skus):
            errors.append("SKU: duplicate inside workbook")
        seen_skus.add(sku)
        if barcode and barcode in seen_barcodes:
            errors.append("Barcode: duplicate inside workbook")
        if barcode:
            seen_barcodes.add(barcode)
        match = existing.get(sku)
        if match and mode == ProductImportMode.CREATE_ONLY:
            errors.append("SKU: already exists; Create New Only never overwrites products")
        if barcode and barcode in existing_barcodes and existing_barcodes[barcode] != sku:
            errors.append("Barcode: already belongs to another product")
        initial = capture(
            "Initial Quantity",
            lambda values=values: _integer(values[11], "Initial Quantity"),
            0,
        )
        reorder = capture(
            "Reorder Level",
            lambda values=values: _integer(values[12], "Reorder Level"),
            0,
        )
        track = capture("Track Stock", lambda values=values: _boolean(values[14], True), True)
        made = capture("Made To Order", lambda values=values: _boolean(values[19], False), False)
        if category == ProductCategory.OPTICAL_SERVICE.value and track:
            errors.append("Track Stock: optical services cannot track stock")
        if initial and not track:
            errors.append("Initial Quantity: requires Track Stock = Yes")
        if match and initial:
            warnings.append("Initial quantity is ignored for an existing product")
        payload = {
            "sku": sku,
            "barcode": barcode,
            "name": name,
            "category": category,
            "brand": capture("Brand", lambda values=values: _text(values[4], 100), None),
            "model": capture("Model", lambda values=values: _text(values[5], 100), None),
            "frame_color": capture("Color", lambda values=values: _text(values[6], 80), None),
            "frame_size": capture("Size", lambda values=values: _text(values[7], 40), None),
            "supplier_id": str(supplier.id) if supplier else None,
            "supplier_name": supplier_name,
            "cost_price_minor": capture("Cost Price", lambda values=values: _money(values[9]), 0),
            "selling_price_minor": capture(
                "Selling Price", lambda values=values: _money(values[10]), 0
            ),
            "initial_quantity": initial,
            "reorder_threshold": reorder,
            "description": capture(
                "Description", lambda values=values: _text(values[13], 4000), None
            ),
            "track_stock": track,
            "frame_material": capture(
                "Frame Material", lambda values=values: _text(values[15], 80), None
            ),
            "frame_shape": capture(
                "Frame Shape", lambda values=values: _text(values[16], 80), None
            ),
            "lens_type": capture("Lens Type", lambda values=values: _text(values[17], 80), None),
            "lens_material": capture(
                "Lens Material", lambda values=values: _text(values[18], 80), None
            ),
            "made_to_order": made,
        }
        if not errors:
            data = {
                key: value
                for key, value in payload.items()
                if key not in {"initial_quantity", "supplier_name"}
            }
            capture("Product", lambda data=data: ProductCreate.model_validate(data))
        rows.append(
            {
                "row": number,
                "data": payload,
                "errors": errors,
                "warnings": warnings,
                "existing": bool(match),
            }
        )
    if not rows:
        raise HTTPException(422, "Products sheet contains no product rows")
    token = secrets.token_urlsafe(32)
    now = datetime.now(UTC)
    valid = sum(not item["errors"] for item in rows)
    preview = ProductImport(
        tenant_id=user.tenant_id,
        actor_id=user.id,
        token_hash=hashlib.sha256(token.encode()).hexdigest(),
        template_version=TEMPLATE_VERSION,
        mode=mode,
        status=ProductImportStatus.PREVIEWED,
        rows_payload=rows,
        total_rows=len(rows),
        valid_rows=valid,
        invalid_rows=len(rows) - valid,
        created_at=now,
        expires_at=now + timedelta(minutes=30),
    )
    db.add(preview)
    db.commit()
    return response(preview, token)


def response(item: ProductImport, token: str | None = None) -> dict:
    result = {
        "id": str(item.id),
        "preview_token": token,
        "template_version": item.template_version,
        "mode": item.mode.value,
        "status": item.status.value,
        "total_rows": item.total_rows,
        "valid_rows": item.valid_rows,
        "invalid_rows": item.invalid_rows,
        "new_products": sum(not row["existing"] and not row["errors"] for row in item.rows_payload),
        "existing_products": sum(row["existing"] for row in item.rows_payload),
        "rows": item.rows_payload,
        "created_count": item.created_count,
        "updated_count": item.updated_count,
        "stock_movements_count": item.stock_movements_count,
        "created_at": item.created_at,
        "completed_at": item.completed_at,
    }
    return result


def confirm_import(db: Session, user: User, token: str, key: str) -> dict:
    digest = hashlib.sha256(token.encode()).hexdigest()
    item = db.scalar(
        select(ProductImport)
        .where(ProductImport.token_hash == digest, ProductImport.tenant_id == user.tenant_id)
        .with_for_update()
    )
    if item is None:
        raise HTTPException(404, "Import preview not found")
    if item.status == ProductImportStatus.COMPLETED:
        return response(item)
    if item.expires_at.replace(tzinfo=UTC) < datetime.now(UTC):
        raise HTTPException(410, "Import preview expired; upload the workbook again")
    if item.invalid_rows:
        raise HTTPException(422, "Correct all invalid rows before confirming the import")
    item.confirmation_key = key
    created = updated = movements = 0
    try:
        for row in item.rows_payload:
            raw = row["data"]
            product = db.scalar(
                select(Product)
                .where(Product.tenant_id == user.tenant_id, Product.sku == raw["sku"])
                .with_for_update()
            )
            supplier_id = UUID(raw["supplier_id"]) if raw["supplier_id"] else None
            raw["supplier_id"] = supplier_id
            if supplier_id and not db.scalar(
                select(Supplier.id).where(
                    Supplier.id == supplier_id,
                    Supplier.tenant_id == user.tenant_id,
                    Supplier.is_archived.is_(False),
                )
            ):
                raise HTTPException(409, f"Row {row['row']}: supplier changed; preview again")
            product_data = {
                key: value
                for key, value in raw.items()
                if key not in {"initial_quantity", "supplier_name"}
            }
            validated = ProductCreate.model_validate(product_data)
            if product:
                if item.mode == ProductImportMode.CREATE_ONLY:
                    raise HTTPException(409, f"Row {row['row']}: SKU now exists; preview again")
                for field, value in validated.model_dump().items():
                    if field != "sku":
                        setattr(product, field, value)
                product.version += 1
                updated += 1
                record_audit(
                    db,
                    tenant_id=user.tenant_id,
                    actor_id=user.id,
                    action="product.import_updated",
                    entity_type="product",
                    entity_id=product.id,
                )
            else:
                product = Product(**validated.model_dump(), tenant_id=user.tenant_id)
                db.add(product)
                db.flush()
                created += 1
                record_audit(
                    db,
                    tenant_id=user.tenant_id,
                    actor_id=user.id,
                    action="product.import_created",
                    entity_type="product",
                    entity_id=product.id,
                )
                if raw["initial_quantity"]:
                    mutate_stock(
                        db,
                        user=user,
                        product=product,
                        data=StockMutation(
                            quantity_delta=raw["initial_quantity"],
                            movement_type=MovementType.OPENING_STOCK,
                            reason="Initial stock from product import",
                            reference=f"IMPORT-{item.id}",
                            expected_version=1,
                        ),
                    )
                    movements += 1
        item.status = ProductImportStatus.COMPLETED
        item.created_count, item.updated_count, item.stock_movements_count = (
            created,
            updated,
            movements,
        )
        item.completed_at = datetime.now(UTC)
        record_audit(
            db,
            tenant_id=user.tenant_id,
            actor_id=user.id,
            action="product_import.completed",
            entity_type="product_import",
            entity_id=item.id,
        )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            409, "Catalog changed during import; preview the workbook again"
        ) from exc
    except Exception:
        db.rollback()
        raise
    return response(item)
